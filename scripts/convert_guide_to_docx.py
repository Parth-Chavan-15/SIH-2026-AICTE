import os
import re
import docx
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

def set_cell_background(cell, fill_hex):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = parse_xml(f'<w:shd {nsdecls("w")} w:fill="{fill_hex}"/>')
    tcPr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = parse_xml(f'<w:tcMar {nsdecls("w")}><w:top w:w="{top}" w:type="dxa"/><w:bottom w:w="{bottom}" w:type="dxa"/><w:left w:w="{left}" w:type="dxa"/><w:right w:w="{right}" w:type="dxa"/></w:tcMar>')
    tcPr.append(tcMar)

def create_styled_document(md_path, docx_path):
    doc = Document()

    # Set page margins (0.75 in)
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.75)
        section.bottom_margin = Inches(0.75)
        section.left_margin = Inches(0.75)
        section.right_margin = Inches(0.75)

    # Configure Default Font
    style_normal = doc.styles['Normal']
    font = style_normal.font
    font.name = 'Calibri'
    font.size = Pt(11)
    font.color.rgb = RGBColor(0x1F, 0x29, 0x37)

    with open(md_path, 'r', encoding='utf-8') as f:
        lines = f.readlines()

    base_dir = os.path.dirname(os.path.abspath(md_path))

    i = 0
    in_code_block = False

    while i < len(lines):
        line = lines[i].rstrip('\r\n')

        # Code block toggle
        if line.startswith('```'):
            in_code_block = not in_code_block
            i += 1
            continue

        if in_code_block:
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Inches(0.3)
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(1)
            run = p.add_run(line)
            run.font.name = 'Consolas'
            run.font.size = Pt(9.5)
            run.font.color.rgb = RGBColor(0x0F, 0x17, 0x2A)
            i += 1
            continue

        # Horizontal Rule
        if line.startswith('---'):
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(8)
            p.paragraph_format.space_after = Pt(8)
            run = p.add_run('_________________________________________________________________________________')
            run.font.color.rgb = RGBColor(0xCB, 0xD5, 0xE1)
            i += 1
            continue

        # Heading 1
        if line.startswith('# '):
            text = line[2:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(18)
            p.paragraph_format.space_after = Pt(6)
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.size = Pt(22)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0A, 0x25, 0x40)
            i += 1
            continue

        # Heading 2
        if line.startswith('## '):
            text = line[3:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(14)
            p.paragraph_format.space_after = Pt(4)
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.size = Pt(15)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x0A, 0x25, 0x40)
            i += 1
            continue

        # Heading 3
        if line.startswith('### '):
            text = line[4:].strip()
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(10)
            p.paragraph_format.space_after = Pt(3)
            run = p.add_run(text)
            run.font.name = 'Calibri'
            run.font.size = Pt(12.5)
            run.font.bold = True
            run.font.color.rgb = RGBColor(0x1E, 0x3A, 0x8A)
            i += 1
            continue

        # Image Embed: ![Caption](./assets/screenshots/...)
        img_match = re.match(r'!\[(.*?)\]\((.*?)\)', line)
        if img_match:
            caption = img_match.group(1)
            img_rel = img_match.group(2)
            img_abs = os.path.normpath(os.path.join(base_dir, img_rel))

            if os.path.exists(img_abs):
                p_img = doc.add_paragraph()
                p_img.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_img.paragraph_format.space_before = Pt(6)
                p_img.paragraph_format.space_after = Pt(4)
                run_img = p_img.add_run()
                run_img.add_picture(img_abs, width=Inches(6.2))

                # Caption
                p_cap = doc.add_paragraph()
                p_cap.alignment = WD_ALIGN_PARAGRAPH.CENTER
                p_cap.paragraph_format.space_before = Pt(0)
                p_cap.paragraph_format.space_after = Pt(8)
                run_cap = p_cap.add_run(f"Figure: {caption}")
                run_cap.font.italic = True
                run_cap.font.size = Pt(9.5)
                run_cap.font.color.rgb = RGBColor(0x64, 0x74, 0x8B)
            i += 1
            continue

        # Tables: line starts with |
        if line.startswith('|') and '|' in line[1:]:
            table_lines = []
            while i < len(lines) and lines[i].strip().startswith('|'):
                table_lines.append(lines[i].strip())
                i += 1

            if len(table_lines) >= 3:
                headers = [c.strip() for c in table_lines[0].strip('|').split('|')]
                data_rows = []
                for tline in table_lines[2:]:
                    cols = [c.strip() for c in tline.strip('|').split('|')]
                    data_rows.append(cols)

                tbl = doc.add_table(rows=len(data_rows) + 1, cols=len(headers))
                tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
                tbl.autofit = True

                # Style header
                hdr_cells = tbl.rows[0].cells
                for ci, h in enumerate(headers):
                    hdr_cells[ci].text = h
                    set_cell_background(hdr_cells[ci], '0A2540')
                    set_cell_margins(hdr_cells[ci], 80, 80, 100, 100)
                    for p in hdr_cells[ci].paragraphs:
                        p.alignment = WD_ALIGN_PARAGRAPH.LEFT
                        for r in p.runs:
                            r.font.bold = True
                            r.font.size = Pt(9.5)
                            r.font.color.rgb = RGBColor(0xFF, 0xFF, 0xFF)

                # Style data rows
                for ri, row_data in enumerate(data_rows):
                    row_cells = tbl.rows[ri + 1].cells
                    bg_color = 'F8FAFC' if ri % 2 == 1 else 'FFFFFF'
                    for ci, ctext in enumerate(row_data):
                        if ci < len(row_cells):
                            # clean basic markdown like ** or `
                            clean_text = ctext.replace('**', '').replace('`', '').replace('<br>', '\n')
                            row_cells[ci].text = clean_text
                            set_cell_background(row_cells[ci], bg_color)
                            set_cell_margins(row_cells[ci], 60, 60, 100, 100)
                            for p in row_cells[ci].paragraphs:
                                for r in p.runs:
                                    r.font.size = Pt(9.5)

                p_after = doc.add_paragraph()
                p_after.paragraph_format.space_before = Pt(4)
                p_after.paragraph_format.space_after = Pt(4)
            continue

        # Bullet lists: line starts with - or * or number
        if line.startswith('- ') or line.startswith('* ') or re.match(r'^\d+\.\s', line):
            p = doc.add_paragraph(style='List Bullet' if not re.match(r'^\d+\.\s', line) else 'List Number')
            p.paragraph_format.space_before = Pt(1)
            p.paragraph_format.space_after = Pt(2)

            text_content = re.sub(r'^[-\*]\s+|\d+\.\s+', '', line)

            # parse bold segments **bold**
            parts = re.split(r'(\*\*.*?\*\*)', text_content)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    p.add_run(part)

            i += 1
            continue

        # Regular Paragraph
        if line.strip():
            p = doc.add_paragraph()
            p.paragraph_format.space_before = Pt(2)
            p.paragraph_format.space_after = Pt(4)

            parts = re.split(r'(\*\*.*?\*\*)', line)
            for part in parts:
                if part.startswith('**') and part.endswith('**'):
                    run = p.add_run(part[2:-2])
                    run.font.bold = True
                else:
                    p.add_run(part)

        i += 1

    doc.save(docx_path)
    print(f"SUCCESS: Document created at {docx_path}")

if __name__ == '__main__':
    md_file = r'c:\Users\parth\Documents\SIH-26-AICTE\docs\NIYAMSETU_PROTOTYPE_WALKTHROUGH_GUIDE.md'
    docx_file = r'c:\Users\parth\Documents\SIH-26-AICTE\docs\NIYAMSETU_PROTOTYPE_WALKTHROUGH_GUIDE.docx'
    create_styled_document(md_file, docx_file)
